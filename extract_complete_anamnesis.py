#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para extrair todas as perguntas/sintomas de anamnese dos arquivos DOCX
e gerar uma estrutura JSON completa para a interface de anamnese.
"""

import json
import re
from pathlib import Path
from docx import Document
from typing import List, Dict, Set
from collections import defaultdict

# Caminhos dos arquivos
BASE_PATH = Path(r"E:\1. Nova pasta\MEUS PROJETOS DE PROGRAMAÇÃO\anamnese-app\src\assets\therapeutic-data")

FILES = [
    "1 da página 001 à 128.docx",
    "2 da página 129 à 238.docx", 
    "3 da página 239 à 240 - EXAMES.docx",
    "4 da página 242 à 424.docx"
]

class AnamnesisExtractor:
    def __init__(self):
        self.questions = []
        self.question_id = 1
        self.categories = set()
        self.conditions = defaultdict(lambda: {"description": "", "symptoms": [], "treatments": []})
        self.protocols = []
        
        # Padrões para identificar perguntas sim/não
        self.yes_no_patterns = [
            r'^[•\-\*]\s*(.+\?)\s*$',  # Perguntas com marcadores
            r'^(.+(?:tem|sente|apresenta|possui|sofre|experimenta).+\?)\s*$',  # Perguntas com verbos típicos
            r'^(.+(?:existe|há|ocorre).+\?)\s*$',
            r'^\d+[\.\)]\s*(.+\?)\s*$',  # Perguntas numeradas
            r'^(.+(?:sim|não|s/n|\( \))\s*(?:sim|não|s/n|\( \))?.*)$',  # Com indicadores sim/não
        ]
        
        # Padrões para sintomas sem pergunta explícita
        self.symptom_patterns = [
            r'^[•\-\*]\s*([A-Z][a-zç]+(?:\s+[a-zç]+)*)\s*$',  # Sintomas listados
            r'^[•\-\*]\s*(.+(?:dor|febre|fadiga|cansaço|insônia|tontura|náusea).+)$',
        ]
        
        # Categorias principais de sistemas do corpo
        self.system_categories = {
            'cardiovascular': ['coração', 'pressão', 'circulação', 'cardiovascular', 'cardíaco'],
            'respiratório': ['pulmão', 'respiração', 'tosse', 'asma', 'bronquite'],
            'digestivo': ['estômago', 'intestino', 'digestão', 'fígado', 'vesícula', 'náusea', 'vômito'],
            'nervoso': ['cérebro', 'nervoso', 'ansiedade', 'depressão', 'stress', 'insônia', 'memória'],
            'endócrino': ['tireoide', 'diabetes', 'hormônio', 'glândula', 'metabolismo'],
            'urinário': ['rim', 'bexiga', 'urina', 'próstata', 'renal'],
            'musculoesquelético': ['osso', 'articulação', 'músculo', 'artrite', 'artrose', 'dor'],
            'pele': ['pele', 'dermatite', 'eczema', 'coceira', 'rash'],
            'reprodutivo': ['menstruação', 'menopausa', 'sexual', 'fertilidade', 'útero', 'ovário'],
            'imunológico': ['imunidade', 'alergia', 'infecção', 'febre'],
            'geral': ['fadiga', 'cansaço', 'peso', 'apetite', 'sono']
        }
        
    def extract_text_from_docx(self, file_path: Path) -> List[str]:
        """Extrai todo o texto do documento DOCX, parágrafo por parágrafo."""
        try:
            doc = Document(file_path)
            paragraphs = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
                    
            # Também extrai texto de tabelas
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            paragraphs.append(text)
                            
            return paragraphs
        except Exception as e:
            print(f"Erro ao ler {file_path}: {e}")
            return []
    
    def categorize_text(self, text: str) -> str:
        """Categoriza o texto baseado em palavras-chave."""
        text_lower = text.lower()
        
        for category, keywords in self.system_categories.items():
            for keyword in keywords:
                if keyword in text_lower:
                    return category
                    
        return 'geral'
    
    def is_question(self, text: str) -> bool:
        """Verifica se o texto é uma pergunta sim/não."""
        # Checa se termina com ?
        if text.endswith('?'):
            return True
            
        # Checa se contém indicadores de sim/não
        if re.search(r'\(\s*\)\s*sim\s*\(\s*\)\s*não', text.lower()):
            return True
            
        if re.search(r'\bsim\b.*\bnão\b|\bnão\b.*\bsim\b', text.lower()):
            return True
            
        return False
    
    def is_symptom(self, text: str) -> bool:
        """Verifica se o texto é um sintoma."""
        symptom_keywords = [
            'dor', 'febre', 'fadiga', 'cansaço', 'insônia', 'tontura', 
            'náusea', 'vômito', 'diarreia', 'constipação', 'tosse',
            'falta de ar', 'palpitação', 'sudorese', 'tremor', 'fraqueza',
            'inchaço', 'edema', 'coceira', 'prurido', 'queimação'
        ]
        
        text_lower = text.lower()
        return any(keyword in text_lower for keyword in symptom_keywords)
    
    def clean_question_text(self, text: str) -> str:
        """Limpa o texto da pergunta removendo marcadores."""
        # Remove marcadores de lista
        text = re.sub(r'^[•\-\*\d+\.\)]\s*', '', text)
        # Remove indicadores sim/não extras
        text = re.sub(r'\(\s*\)\s*(?:sim|não)\s*\(\s*\)\s*(?:sim|não)\s*', '', text, flags=re.IGNORECASE)
        return text.strip()
    
    def extract_condition_info(self, text: str) -> Dict:
        """Extrai informações sobre condições/diagnósticos."""
        # Padrões para identificar condições
        condition_indicators = [
            r'diagnóstico\s*:\s*(.+)',
            r'condição\s*:\s*(.+)',
            r'patologia\s*:\s*(.+)',
            r'doença\s*:\s*(.+)',
        ]
        
        for pattern in condition_indicators:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return {
                    'name': match.group(1).strip(),
                    'description': text
                }
        
        return None
    
    def extract_treatment_info(self, text: str) -> Dict:
        """Extrai informações sobre tratamentos."""
        treatment_indicators = [
            r'tratamento\s*:\s*(.+)',
            r'terapia\s*:\s*(.+)',
            r'protocolo\s*:\s*(.+)',
            r'recomendação\s*:\s*(.+)',
            r'prescrição\s*:\s*(.+)',
        ]
        
        for pattern in treatment_indicators:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                return {
                    'type': 'treatment',
                    'description': text,
                    'details': match.group(1).strip()
                }
        
        return None
    
    def extract_protocol_info(self, text: str) -> Dict:
        """Extrai informações sobre protocolos."""
        if 'protocolo' in text.lower():
            return {
                'name': text,
                'description': text,
                'steps': []
            }
        return None
    
    def process_paragraphs(self, paragraphs: List[str], source_file: str):
        """Processa os parágrafos extraídos."""
        current_condition = None
        
        for i, text in enumerate(paragraphs):
            # Pula textos muito curtos ou vazios
            if len(text) < 3:
                continue
            
            # Identifica perguntas
            if self.is_question(text):
                category = self.categorize_text(text)
                self.categories.add(category)
                
                question = {
                    'id': f'q{self.question_id}',
                    'text': self.clean_question_text(text),
                    'category': category,
                    'type': 'yes_no',
                    'source': source_file,
                    'relatedConditions': []
                }
                
                # Procura condições relacionadas nos parágrafos próximos
                context_range = range(max(0, i-2), min(len(paragraphs), i+3))
                for j in context_range:
                    if j != i:
                        condition_info = self.extract_condition_info(paragraphs[j])
                        if condition_info and condition_info['name'] not in question['relatedConditions']:
                            question['relatedConditions'].append(condition_info['name'])
                            self.conditions[condition_info['name']]['description'] = condition_info['description']
                
                self.questions.append(question)
                self.question_id += 1
            
            # Identifica sintomas que não são perguntas
            elif self.is_symptom(text) and not text.endswith('?'):
                category = self.categorize_text(text)
                self.categories.add(category)
                
                # Converte sintoma em pergunta
                symptom_text = self.clean_question_text(text)
                if not symptom_text.startswith(('Você', 'Tem', 'Sente', 'Apresenta')):
                    symptom_text = f"Você apresenta {symptom_text.lower()}?"
                
                question = {
                    'id': f'q{self.question_id}',
                    'text': symptom_text,
                    'category': category,
                    'type': 'yes_no',
                    'source': source_file,
                    'relatedConditions': []
                }
                
                self.questions.append(question)
                self.question_id += 1
            
            # Identifica condições/diagnósticos
            condition_info = self.extract_condition_info(text)
            if condition_info:
                current_condition = condition_info['name']
                self.conditions[current_condition]['description'] = condition_info['description']
            
            # Identifica tratamentos
            treatment_info = self.extract_treatment_info(text)
            if treatment_info and current_condition:
                self.conditions[current_condition]['treatments'].append(treatment_info)
            
            # Identifica protocolos
            protocol_info = self.extract_protocol_info(text)
            if protocol_info:
                self.protocols.append(protocol_info)
    
    def process_files(self):
        """Processa todos os arquivos DOCX."""
        for file_name in FILES:
            file_path = BASE_PATH / file_name
            print(f"\nProcessando: {file_name}")
            
            if not file_path.exists():
                print(f"  ⚠️  Arquivo não encontrado: {file_path}")
                continue
            
            paragraphs = self.extract_text_from_docx(file_path)
            print(f"  📄 {len(paragraphs)} parágrafos extraídos")
            
            self.process_paragraphs(paragraphs, file_name)
            print(f"  ✅ Processado - Total de perguntas até agora: {len(self.questions)}")
    
    def generate_json_structure(self) -> Dict:
        """Gera a estrutura JSON final."""
        return {
            'metadata': {
                'version': '1.0',
                'extractionDate': '2026-02-03',
                'totalQuestions': len(self.questions),
                'totalCategories': len(self.categories),
                'totalConditions': len(self.conditions),
                'totalProtocols': len(self.protocols),
                'sourceFiles': FILES
            },
            'categories': sorted(list(self.categories)),
            'questions': self.questions,
            'conditions': dict(self.conditions),
            'protocols': self.protocols,
            'categoriesDescription': {
                'cardiovascular': 'Sistema Cardiovascular',
                'respiratório': 'Sistema Respiratório',
                'digestivo': 'Sistema Digestivo',
                'nervoso': 'Sistema Nervoso',
                'endócrino': 'Sistema Endócrino',
                'urinário': 'Sistema Urinário',
                'musculoesquelético': 'Sistema Musculoesquelético',
                'pele': 'Pele e Anexos',
                'reprodutivo': 'Sistema Reprodutivo',
                'imunológico': 'Sistema Imunológico',
                'geral': 'Sintomas Gerais'
            }
        }
    
    def save_to_file(self, output_file: str):
        """Salva a estrutura JSON em arquivo."""
        json_data = self.generate_json_structure()
        
        output_path = Path(output_file)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(json_data, f, ensure_ascii=False, indent=2)
        
        print(f"\n{'='*70}")
        print(f"✅ Extração concluída!")
        print(f"{'='*70}")
        print(f"📊 Estatísticas:")
        print(f"  • Total de perguntas: {len(self.questions)}")
        print(f"  • Total de categorias: {len(self.categories)}")
        print(f"  • Total de condições: {len(self.conditions)}")
        print(f"  • Total de protocolos: {len(self.protocols)}")
        print(f"\n💾 Arquivo salvo: {output_path}")
        print(f"{'='*70}\n")
        
        # Mostra amostra de perguntas por categoria
        print("📋 Amostra de perguntas por categoria:")
        for category in sorted(self.categories):
            questions_in_category = [q for q in self.questions if q['category'] == category]
            print(f"\n  {category.upper()} ({len(questions_in_category)} perguntas)")
            for q in questions_in_category[:3]:  # Mostra 3 exemplos
                print(f"    - {q['text'][:80]}...")


def main():
    print("="*70)
    print("EXTRAÇÃO DE ANAMNESE - ARQUIVOS DOCX")
    print("="*70)
    
    extractor = AnamnesisExtractor()
    extractor.process_files()
    
    # Salva o resultado
    output_file = r"E:\1. Nova pasta\MEUS PROJETOS DE PROGRAMAÇÃO\anamnese-app\anamnesis_complete_data.json"
    extractor.save_to_file(output_file)


if __name__ == "__main__":
    main()
