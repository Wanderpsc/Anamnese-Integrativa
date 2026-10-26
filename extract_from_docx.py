from docx import Document
import json
import re

docx_path = r"E:\1. Nova pasta\MEUS PROJETOS DE PROGRAMAÇÃO\anamnese-app\src\assets\therapeutic-data\4 da página 242 à 424.docx"

try:
    doc = Document(docx_path)
    
    print(f"Total de parágrafos: {len(doc.paragraphs)}")
    
    # Extrair todo o texto
    all_text = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            all_text.append(text)
    
    # Procurar padrões de exames
    exams = []
    current_exam = {}
    
    for i, text in enumerate(all_text):
        # Detectar início de novo exame (geralmente tem padrão de título)
        if re.search(r'^[A-Z][a-zá-ú]+\s*\([A-Z][a-z]*\)', text) or \
           re.search(r'^(Hemácias|Hemoglobina|Hematócrito|VCM|HCM|CHCM|RDW|Leucócitos|Neutrófilos|Linfócitos|Monócitos|Eosinófilos|Basófilos|Plaquetas)', text, re.IGNORECASE):
            
            if current_exam:
                exams.append(current_exam)
            
            current_exam = {
                'nome': text,
                'conteudo': [],
                'valores_referencia': None,
                'valores_ideais': None
            }
        
        elif current_exam:
            current_exam['conteudo'].append(text)
            
            # Capturar valores de referência
            if 'valores de referência' in text.lower() or 'valores laboratoriais' in text.lower():
                # Próximas linhas devem conter os valores
                if i+1 < len(all_text):
                    current_exam['valores_referencia'] = all_text[i+1]
            
            # Capturar valores ideais
            if 'valor ideal funcional' in text.lower() or 'valores funcionais ideais' in text.lower():
                if i+1 < len(all_text):
                    current_exam['valores_ideais'] = all_text[i+1]
    
    # Adicionar último exame
    if current_exam:
        exams.append(current_exam)
    
    print(f"\n\nTotal de exames encontrados: {len(exams)}")
    print("\n" + "="*80)
    print("PRIMEIROS 10 EXAMES ESTRUTURADOS:")
    print("="*80)
    
    for i, exam in enumerate(exams[:10]):
        print(f"\n{'='*80}")
        print(f"EXAME {i+1}: {exam['nome']}")
        print('='*80)
        print(f"Valores de Referência: {exam.get('valores_referencia', 'Não encontrado')}")
        print(f"Valores Ideais: {exam.get('valores_ideais', 'Não encontrado')}")
        print(f"Conteúdo (primeiras 5 linhas):")
        for line in exam['conteudo'][:5]:
            print(f"  - {line}")
    
    # Salvar JSON
    with open('exames_estruturados.json', 'w', encoding='utf-8') as f:
        json.dump(exams, f, ensure_ascii=False, indent=2)
    
    print(f"\n\nArquivo 'exames_estruturados.json' criado com {len(exams)} exames")
            
except Exception as e:
    print(f"Erro ao ler DOCX: {e}")
    import traceback
    traceback.print_exc()
